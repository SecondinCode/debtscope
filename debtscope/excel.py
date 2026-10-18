"""Excel report: live formulas, maturity-wall chart and a full audit trail.

Conventions (standard sell-side modelling):
    blue font  = data extracted from the filing (hardcoded input)
    black font = formula
    amounts    = reporting currency, millions
Every figure on the Summary tab is a formula pointing at the XBRL or
Instruments tabs, so editing an input (e.g. fixing a principal) flows through.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .analytics import principal_universe
from .models import DebtInstrument, DebtProfile
from .xbrl import AFTER_FIVE

NAVY = "1F2A44"
SLATE = "5B6B82"
LIGHT = "EEF2F7"
BLUE_INPUT = "0000FF"
FONT = "Arial"

TITLE = Font(name=FONT, size=18, bold=True, color=NAVY)
SUB = Font(name=FONT, size=10, color=SLATE)
H = Font(name=FONT, size=11, bold=True, color="FFFFFF")
LBL = Font(name=FONT, size=10, color="222222")
INP = Font(name=FONT, size=10, color=BLUE_INPUT)
CALC = Font(name=FONT, size=10, color="000000")
BOLD = Font(name=FONT, size=10, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color=SLATE)
WARN = Font(name=FONT, size=10, color="9C2A00")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor=LIGHT)
THIN = Side(style="thin", color="C9D1DC")
BOX = Border(bottom=THIN)

MM = '#,##0;(#,##0);"-"'
MM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT3 = '0.000%;(0.000%);"-"'
MULT = '0.0"x";(0.0"x");"-"'
YRS = '0.0" yrs";;"-"'
DATE = "dd-mmm-yyyy"

SNAPSHOT_LABELS = {
    "total_debt": "Total term debt (carrying)",
    "debt_noncurrent": "Term debt - non-current",
    "debt_current": "Term debt - current",
    "commercial_paper": "Commercial paper",
    "face_amount": "Debt face amount",
    "cash": "Cash & cash equivalents",
    "marketable_securities_current": "Current marketable securities",
    "operating_income": "Operating income (FY)",
    "interest_expense": "Interest expense (FY)",
    "interest_paid": "Cash interest paid (FY)",
    "depreciation": "D&A (FY)",
    "shares_outstanding": "Shares outstanding (mm)",
}


def _x(key: str) -> str:
    """Formula fragment: value for `key` from the XBRL tab, #N/A when missing
    (so IFERROR catches it; a text "n/a" would leak into sums as #VALUE!)."""
    v = f'INDEX(XBRL!$C:$C,MATCH("{key}",XBRL!$A:$A,0))'
    return f"IF(ISNUMBER({v}),{v},NA())"


def _header(ws, row: int, text: str, width: int = 4) -> None:
    ws.cell(row=row, column=1, value=text).font = H
    for c in range(1, width + 1):
        ws.cell(row=row, column=c).fill = HEAD_FILL
    ws.row_dimensions[row].height = 20


def _widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _maturity_date(i: DebtInstrument) -> tuple[date | None, bool]:
    if not i.maturity_year or i.is_aggregate:
        return None, False
    est = i.maturity_month is None or i.maturity_day is None
    return date(i.maturity_year, i.maturity_month or 7, i.maturity_day or 1), est


# ------------------------------------------------------------------ sheets
def _xbrl_sheet(wb: Workbook, p: DebtProfile) -> None:
    ws = wb.create_sheet("XBRL")
    ws["A1"] = "XBRL facts reported in this filing"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    ws["A2"] = ("Source: SEC companyfacts API, filtered on accession and period end. "
                f"Values in {p.reporting_currency} millions.")
    ws["A2"].font = NOTE
    heads = ["Key", "Description", "Value (mm)", "XBRL tag", "Period end", "Accession"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = H
        cell.fill = HEAD_FILL
    r = 5
    for key, label in SNAPSHOT_LABELS.items():
        fact = p.xbrl.get(key)
        ws.cell(row=r, column=1, value=key).font = LBL
        ws.cell(row=r, column=2, value=label).font = LBL
        if fact:
            ws.cell(row=r, column=3, value=fact.value / 1e6).font = INP
            ws.cell(row=r, column=4, value=fact.tag).font = NOTE
            ws.cell(row=r, column=5, value=fact.period_end).font = NOTE
            ws.cell(row=r, column=6, value=fact.accession).font = NOTE
        else:
            ws.cell(row=r, column=3, value="n/a").font = NOTE
            ws.cell(row=r, column=4, value="not reported under any mapped tag").font = NOTE
        ws.cell(row=r, column=3).number_format = MM1
        r += 1
    r += 1
    _header(ws, r, "Contractual maturity schedule (principal)", 6)
    r += 1
    sched_first = r
    names = {0: "Remainder of fiscal year", 1: "Year 1 (next 12 months)", 2: "Year 2",
             3: "Year 3", 4: "Year 4", 5: "Year 5", AFTER_FIVE: "Thereafter"}
    for bucket, name in names.items():
        fact = p.maturity_schedule.get(bucket)
        if not fact and bucket == 0:
            continue
        ws.cell(row=r, column=1, value=f"sched_{bucket}").font = LBL
        ws.cell(row=r, column=2, value=name).font = LBL
        if fact:
            ws.cell(row=r, column=3, value=fact.value / 1e6).font = INP
            ws.cell(row=r, column=4, value=fact.tag.split(":")[-1]).font = NOTE
            ws.cell(row=r, column=5, value=fact.period_end).font = NOTE
            ws.cell(row=r, column=6, value=fact.accession).font = NOTE
        else:
            ws.cell(row=r, column=3, value="n/a").font = NOTE
        ws.cell(row=r, column=3).number_format = MM1
        r += 1
    ws.cell(row=r, column=1, value="sched_total").font = BOLD
    ws.cell(row=r, column=2, value="Total scheduled principal").font = BOLD
    ws.cell(row=r, column=3, value=f"=SUM(C{sched_first}:C{r - 1})").font = BOLD
    ws.cell(row=r, column=3).number_format = MM1
    _widths(ws, {"A": 22, "B": 32, "C": 14, "D": 62, "E": 12, "F": 24})
    ws.freeze_panes = "A5"


def _instruments_sheet(wb: Workbook, p: DebtProfile) -> int:
    ws = wb.create_sheet("Instruments")
    ws["A1"] = "Debt instruments extracted from the filing"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    ws["A2"] = (f"Principal in {p.reporting_currency} millions (reporting currency, as in the footnote). "
                "Denomination '*' = assumed. 'In totals' = No when the note is already inside an "
                "aggregated row, redeemed, or has no disclosed principal (avoids double counting).")
    ws["A2"].font = NOTE
    heads = ["#", "Instrument", "Coupon", "Rate type", "Maturity", "Date est.",
             "Years to maturity", "Principal (mm)", "Denomination", "Status", "In totals",
             "Symbol", "Exchange", "Sources", "Flags", "Maturity year", "Effective rate"]
    hr = 4
    for c, h in enumerate(heads, 1):
        ws.cell(row=hr, column=c, value=h)
    counted = {id(i) for i in principal_universe(p.instruments)}
    r = hr
    for n, inst in enumerate(p.instruments, 1):
        r = hr + n
        mdate, est = _maturity_date(inst)
        rate_type = "Floating" if inst.floating else ("Fixed (range)" if inst.is_aggregate else "Fixed")
        vals = [
            n, inst.label,
            None if inst.is_aggregate or inst.coupon is None else inst.coupon / 100,
            rate_type, mdate, "Yes" if est else ("" if mdate is None else "No"),
            # YEARFRAC ignores date order, so past maturities are zeroed first.
            f"=IF(E{r}=\"\",0,IF(E{r}<=Summary!$C$8,0,YEARFRAC(Summary!$C$8,E{r},1)))",
            None if inst.amount is None else inst.amount / 1e6,
            inst.currency, inst.status, "Yes" if id(inst) in counted else "No",
            inst.symbol, inst.exchange, ", ".join(inst.sources), "; ".join(inst.flags),
            inst.maturity_year_to or inst.maturity_year,
            None if inst.effective_rate is None else inst.effective_rate / 100,
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = CALC if c == 7 else INP if c in (3, 5, 8) else LBL
            cell.alignment = Alignment(vertical="top", wrap_text=c in (2, 15))
        ws.cell(row=r, column=3).number_format = PCT3
        ws.cell(row=r, column=5).number_format = DATE
        ws.cell(row=r, column=7).number_format = YRS
        ws.cell(row=r, column=8).number_format = MM
        ws.cell(row=r, column=17).number_format = PCT3
        if inst.effective_rate is None and inst.effective_rate_text:
            ws.cell(row=r, column=17).comment = Comment(
                f"Filing gives a range: {inst.effective_rate_text}", "DebtScope")
        if inst.flags:
            ws.cell(row=r, column=2).comment = Comment("\n".join(inst.flags), "DebtScope")
    last = max(r, hr + 1)
    if not p.instruments:
        ws.cell(row=hr + 1, column=2, value="No debt instruments found in the document").font = WARN
    ref = f"A{hr}:{get_column_letter(len(heads))}{last}"
    table = Table(displayName="Instruments", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    if p.instruments:
        ws.conditional_formatting.add(
            f"H{hr + 1}:H{last}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                              color="5B8DEF", showValue=True))
        ws.conditional_formatting.add(
            f"G{hr + 1}:G{last}", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                                 mid_type="num", mid_value=5, mid_color="FFEB84",
                                                 end_type="num", end_value=15, end_color="63BE7B"))
        ws.conditional_formatting.add(
            f"J{hr + 1}:J{last}", CellIsRule(operator="notEqual", formula=['"outstanding"'],
                                             font=Font(name=FONT, color="9C2A00", bold=True)))
    _widths(ws, {"A": 5, "B": 38, "C": 10, "D": 13, "E": 13, "F": 9, "G": 15, "H": 15,
                 "I": 13, "J": 17, "K": 10, "L": 10, "M": 26, "N": 34, "O": 60, "P": 13, "Q": 13})
    ws.freeze_panes = "C5"
    return last


def _maturity_sheet(wb: Workbook, p: DebtProfile, inst_last: int) -> dict:
    """Two views, two charts.

    A. Contractual maturities from XBRL (years 1-5 + thereafter): audited and
       complete, the right basis for refinancing risk.
    B. Bond-level profile by calendar year: principal when the filing gives
       it bond by bond, otherwise the number of notes maturing each year.
    """
    ws = wb.create_sheet("Maturity Wall")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Maturity wall"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    ws["A2"] = ("A: contractual principal by fiscal year after the period end (XBRL, complete). "
                "B: individual notes by calendar year of maturity.")
    ws["A2"].font = NOTE
    fy = int(p.report_date[:4])
    ccy = p.reporting_currency

    # ---------------- A. XBRL schedule
    hr = 4
    for c, h in enumerate(["A. Contractual maturities (XBRL)", "Fiscal year",
                           f"Principal ({ccy} mm)", "% of total", "Cumulative"], 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = H
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center", wrap_text=True)
    buckets = ([0] if 0 in p.maturity_schedule else []) + [1, 2, 3, 4, 5, AFTER_FIVE]
    names = {0: "Remainder of fiscal year", 1: "Year 1", 2: "Year 2", 3: "Year 3",
             4: "Year 4", 5: "Year 5", AFTER_FIVE: "Thereafter"}
    first = hr + 1
    last = first + len(buckets) - 1
    total_row = last + 1
    rows: dict[int, int] = {}
    for k, b in enumerate(buckets):
        r = first + k
        rows[b] = r
        fy_label = (f"FY{fy}" if b == 0 else f"After FY{fy + 5}" if b == AFTER_FIVE else f"FY{fy + b}")
        ws.cell(row=r, column=1, value=names[b]).font = LBL
        ws.cell(row=r, column=2, value=fy_label).font = LBL
        ws.cell(row=r, column=3, value=f'=IFERROR({_x(f"sched_{b}")},0)').font = CALC
        ws.cell(row=r, column=4, value=f"=IFERROR(C{r}/$C${total_row},0)").font = CALC
        ws.cell(row=r, column=5, value=f"=IFERROR(SUM($C${first}:C{r})/$C${total_row},0)").font = CALC
        ws.cell(row=r, column=3).number_format = MM
        ws.cell(row=r, column=4).number_format = PCT
        ws.cell(row=r, column=5).number_format = PCT
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        if k % 2:
            for c in range(1, 6):
                ws.cell(row=r, column=c).fill = BAND_FILL
    ws.cell(row=total_row, column=1, value="Total scheduled principal").font = BOLD
    ws.cell(row=total_row, column=3, value=f"=SUM(C{first}:C{last})").font = BOLD
    ws.cell(row=total_row, column=3).number_format = MM
    for c in range(1, 6):
        ws.cell(row=total_row, column=c).border = Border(top=Side(style="thin", color=NAVY))
    if not p.maturity_schedule:
        ws.cell(row=total_row + 1, column=1,
                value="No maturity schedule tagged in XBRL for this filing.").font = WARN

    chart = BarChart()
    chart.type = "col"
    chart.title = f"Contractual maturities ({ccy} mm)"
    chart.y_axis.numFmt = "#,##0"
    chart.y_axis.majorGridlines = None
    chart.add_data(Reference(ws, min_col=3, min_row=hr, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=first, max_row=last))
    chart.series[0].graphicalProperties.solidFill = "1F4E9A"
    chart.gapWidth = 60
    line = LineChart()
    line.add_data(Reference(ws, min_col=5, min_row=hr, max_row=last), titles_from_data=True)
    line.y_axis.axId = 200
    line.y_axis.numFmt = "0%"
    line.y_axis.scaling.min, line.y_axis.scaling.max = 0, 1
    line.y_axis.crosses = "max"
    line.y_axis.majorGridlines = None
    line.series[0].graphicalProperties.line.solidFill = "E07A1F"
    line.series[0].graphicalProperties.line.width = 22000
    line.series[0].smooth = False
    for axis in (chart.x_axis, chart.y_axis, line.y_axis):
        axis.delete = False
    chart += line
    chart.legend.position = "b"
    chart.height, chart.width = 8, 18
    ws.add_chart(chart, "H4")

    # ---------------- B. Bond-level profile
    live = [i for i in p.instruments if not i.is_aggregate and i.maturity_year
            and i.status in ("outstanding", "listed in exhibits")]
    priced = sum(i.amount for i in live if i.amount and i.status == "outstanding")
    total_debt = p.metrics.get("xbrl_total_debt") or 0
    principal_mode = bool(total_debt) and priced / total_debt >= 0.6
    b_hr = max(total_row + 4, 22)
    for c, h in enumerate(["B. Individual notes", f"Principal ({ccy} mm)", "# notes"], 1):
        cell = ws.cell(row=b_hr, column=c, value=h)
        cell.font = H
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center", wrap_text=True)
    ws.cell(row=b_hr - 1, column=1, value=(
        "Principal disclosed bond by bond: chart shows amounts." if principal_mode else
        "Amounts are not disclosed bond by bond (aggregated rows or exhibit list only): "
        "chart shows how many notes mature each year.")).font = NOTE
    years = sorted({i.maturity_year for i in live})
    b_first = b_hr + 1
    rng_amt = f"Instruments!$H$5:$H${inst_last}"
    rng_yr = f"Instruments!$P$5:$P${inst_last}"
    rng_in = f"Instruments!$K$5:$K${inst_last}"
    rng_st = f"Instruments!$J$5:$J${inst_last}"
    rng_ag = f"Instruments!$D$5:$D${inst_last}"
    if years:
        span = list(range(max(min(years), fy), min(max(years), fy + 45) + 1))
    else:
        span = []
    for k, y in enumerate(span):
        r = b_first + k
        ws.cell(row=r, column=1, value=y).number_format = "0"
        ws.cell(row=r, column=2, value=(
            f'=SUMIFS({rng_amt},{rng_yr},A{r},{rng_in},"Yes",{rng_ag},"<>Fixed (range)")'))
        ws.cell(row=r, column=3, value=(
            f'=COUNTIFS({rng_yr},A{r},{rng_st},"outstanding",{rng_ag},"<>Fixed (range)")'
            f'+COUNTIFS({rng_yr},A{r},{rng_st},"listed in exhibits")'))
        for c in (1, 2, 3):
            ws.cell(row=r, column=c).font = CALC
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2).number_format = MM
        if k % 2:
            for c in (1, 2, 3):
                ws.cell(row=r, column=c).fill = BAND_FILL
    b_last = b_first + len(span) - 1
    if span:
        bars = BarChart()
        bars.type = "col"
        col = 2 if principal_mode else 3
        bars.title = (f"Notes by maturity year ({ccy} mm)" if principal_mode
                      else "Number of notes maturing per year")
        bars.add_data(Reference(ws, min_col=col, min_row=b_hr, max_row=b_last), titles_from_data=True)
        bars.set_categories(Reference(ws, min_col=1, min_row=b_first, max_row=b_last))
        bars.series[0].graphicalProperties.solidFill = "5B8DEF"
        bars.y_axis.majorGridlines = None
        bars.y_axis.numFmt = "#,##0" if principal_mode else "0"
        bars.gapWidth = 40
        bars.legend = None
        for axis in (bars.x_axis, bars.y_axis):
            axis.delete = False
        bars.height, bars.width = 8, 26
        ws.add_chart(bars, f"H{b_hr}")
    else:
        ws.cell(row=b_first, column=1, value="No individual notes identified.").font = NOTE
    _widths(ws, {"A": 30, "B": 16, "C": 16, "D": 12, "E": 12})
    return {"y1": rows[1], "y5": rows[5], "b_first": b_first, "b_last": max(b_first, b_last),
            "principal_mode": principal_mode}


def _audit_sheet(wb: Workbook, p: DebtProfile) -> None:
    ws = wb.create_sheet("Audit Trail")
    ws["A1"] = "Evidence for every extracted instrument"
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
    ws["A2"] = f"Filing: {p.url}"
    ws["A2"].font = NOTE
    heads = ["#", "Instrument", "Source", "Context", "Amount (mm)", "Currency", "Extract from filing"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = H
        cell.fill = HEAD_FILL
    r = 5
    for n, inst in enumerate(p.instruments, 1):
        for ev in inst.evidence:
            vals = [n, inst.label, ev.source, ev.note,
                    None if ev.amount is None else ev.amount / 1e6, ev.currency, ev.snippet]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.font = LBL if c != 7 else NOTE
                cell.alignment = Alignment(vertical="top", wrap_text=c in (4, 7))
            ws.cell(row=r, column=5).number_format = MM
            r += 1
    _widths(ws, {"A": 5, "B": 32, "C": 15, "D": 28, "E": 13, "F": 10, "G": 100})
    ws.freeze_panes = "A5"


def _summary_sheet(ws, p: DebtProfile, inst_last: int, mw: dict) -> None:
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ccy = p.reporting_currency
    ws["A1"] = f"{p.company} ({p.ticker})" if p.ticker else p.company
    ws["A1"].font = TITLE
    ws["A2"] = "Debt structure & maturity profile - generated by DebtScope from SEC EDGAR"
    ws["A2"].font = SUB

    _header(ws, 4, "Filing")
    rows = [
        ("Company", p.company), ("Ticker / CIK", f"{p.ticker} / {p.cik}"),
        ("Form / accession", f"{p.form} / {p.accession}"),
        ("Period end", date.fromisoformat(p.report_date)),
        ("Filed", date.fromisoformat(p.filing_date)),
        ("Units", f"{ccy} millions (document scale: {p.scale_label})"),
        ("Source document", p.url),
    ]
    for k, (label, val) in enumerate(rows, 5):
        ws.cell(row=k, column=1, value=label).font = LBL
        cell = ws.cell(row=k, column=3, value=val)
        cell.font = INP
        if isinstance(val, date):
            cell.number_format = DATE
            cell.alignment = Alignment(horizontal="left")
    ws["C11"].hyperlink = p.url
    ws["C11"].font = Font(name=FONT, size=10, color="1F4E9A", underline="single")
    ws["D8"] = "Drives all years-to-maturity formulas"
    ws["D8"].font = NOTE

    il = inst_last
    J, D, H, K, C, G, Q = (f"Instruments!$J$5:$J${il}", f"Instruments!$D$5:$D${il}",
                           f"Instruments!$H$5:$H${il}", f"Instruments!$K$5:$K${il}",
                           f"Instruments!$C$5:$C${il}", f"Instruments!$G$5:$G${il}",
                           f"Instruments!$Q$5:$Q${il}")
    peak_col = "B" if mw["principal_mode"] else "C"
    bf, bl = mw["b_first"], mw["b_last"]
    sections = [
        ("Balance sheet & credit (XBRL, %s mm)" % ccy, [
            ("Total term debt", f"=IFERROR({_x('total_debt')},\"n/a\")", MM, "LongTermDebt (carrying value)"),
            ("Commercial paper", f"=IFERROR({_x('commercial_paper')},0)", MM, "Short-term funding, outside term debt"),
            ("Gross debt", "=IFERROR(C14+C15,\"n/a\")", MM, "Term debt + commercial paper"),
            ("Cash & current marketable securities",
             f"=IFERROR({_x('cash')},0)+IFERROR({_x('marketable_securities_current')},0)", MM,
             "Non-current securities excluded (conservative)"),
            ("Net debt", "=IFERROR(C16-C17,\"n/a\")", MM, "Negative = net cash"),
            ("EBITDA (operating income + D&A)",
             f"=IFERROR({_x('operating_income')}+{_x('depreciation')},\"n/a\")", MM, "Unadjusted, FY"),
            ("Net debt / EBITDA", "=IFERROR(C18/C19,\"n/a\")", MULT, "Leverage"),
            ("Interest coverage (EBIT / interest)",
             f"=IFERROR({_x('operating_income')}/IFERROR({_x('interest_expense')},{_x('interest_paid')}),\"n/a\")", MULT,
             "Falls back to cash interest paid if expense is not tagged"),
            ("Implied cost of debt", f"=IFERROR(IFERROR({_x('interest_expense')},{_x('interest_paid')})/C16,\"n/a\")", PCT,
             "Interest / period-end gross debt"),
        ]),
        ("Bond-level extraction", [
            ("Instruments identified", f"=COUNTA(Instruments!$B$5:$B${il})", "0", ""),
            ("Outstanding", f'=COUNTIF({J},"outstanding")', "0", ""),
            ("Extracted principal (no double count)", f'=SUMIFS({H},{K},"Yes")', MM, ""),
            ("Coverage vs reported principal",
             f"=IFERROR(C27/{_x('sched_total')},IFERROR(C27/C14,\"n/a\"))", PCT,
             "Face vs XBRL maturity schedule (else vs carrying value)"),
            ("Weighted avg coupon (fixed, outstanding)",
             f'=IFERROR(SUMPRODUCT(({J}="outstanding")*({D}="Fixed")*{C}*{H})/'
             f'SUMPRODUCT(({J}="outstanding")*({D}="Fixed")*{H}),"n/a")', "0.00%", "Principal-weighted"),
            ("Weighted avg maturity", f'=IFERROR(SUMPRODUCT(({J}="outstanding")*({G}>0)*{G}*{H})/'
             f'SUMPRODUCT(({J}="outstanding")*({G}>0)*{H}),"n/a")', YRS, "Bond-level, from period end"),
            ("Weighted avg effective rate (aggregated rows)",
             f'=IFERROR(SUMPRODUCT(({D}="Fixed (range)")*({J}="outstanding")*{Q}*{H})/'
             f'SUMPRODUCT(({D}="Fixed (range)")*({J}="outstanding")*({Q}>0)*{H}),"n/a")', "0.00%",
             "When the filing reports notes only in aggregate"),
            ("Notes listed only in exhibits", f'=COUNTIF({J},"listed in exhibits")', "0",
             "Issued; current balance not disclosed bond by bond"),
        ]),
        ("Refinancing risk (XBRL schedule)", [
            ("Due within 24 months", f"=IFERROR(({_x('sched_1')}+{_x('sched_2')})/{_x('sched_total')},\"n/a\")",
             PCT, "Share of scheduled principal"),
            ("Largest single-year bucket",
             f"=IFERROR(MAX('Maturity Wall'!$C${mw['y1']}:$C${mw['y5']})/{_x('sched_total')},\"n/a\")",
             PCT, "Concentration in years 1-5"),
            ("Peak maturity year (individual notes)",
             f"=IFERROR(INDEX('Maturity Wall'!$A${bf}:$A${bl},MATCH(MAX('Maturity Wall'!${peak_col}${bf}:${peak_col}${bl}),"
             f"'Maturity Wall'!${peak_col}${bf}:${peak_col}${bl},0)),\"n/a\")", "0",
             "By principal" if mw["principal_mode"] else "By number of notes (amounts not disclosed)"),
        ]),
    ]
    r = 13
    for title, items in sections:
        _header(ws, r, title)
        r += 1
        for label, formula, fmt, note in items:
            ws.cell(row=r, column=1, value=label).font = LBL
            cell = ws.cell(row=r, column=3, value=formula)
            cell.font = CALC
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
            ws.cell(row=r, column=4, value=note).font = NOTE
            for c in range(1, 5):
                ws.cell(row=r, column=c).border = BOX
            r += 1
        r += 1
    # Rows referenced above (C14..C19, C27) must match this layout.
    assert ws["A14"].value == "Total term debt" and ws["A27"].value.startswith("Extracted principal")

    _header(ws, r, "Data-quality warnings")
    r += 1
    for w in p.warnings or ["None raised - still spot-check the Audit Trail tab."]:
        ws.cell(row=r, column=1, value="- " + w).font = WARN if p.warnings else NOTE
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=(
        "Automated extraction from public SEC filings. Verify against the filing before use; "
        "not investment advice.")).font = NOTE
    _widths(ws, {"A": 40, "B": 2, "C": 22, "D": 52})


def write_workbook(p: DebtProfile, path: str | Path) -> Path:
    path = Path(path)
    wb = Workbook()
    summary = wb.active
    _xbrl_sheet(wb, p)
    inst_last = _instruments_sheet(wb, p)
    mw = _maturity_sheet(wb, p, inst_last)
    _audit_sheet(wb, p)
    _summary_sheet(summary, p, inst_last, mw)
    wb._sheets = [wb["Summary"], wb["Maturity Wall"], wb["Instruments"], wb["XBRL"], wb["Audit Trail"]]
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # Excel computes every formula on open
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title == "Summary" else SLATE
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# -------------------------------------------------------------- comparison
COMPARE_ROWS = [
    ("Period end", lambda p: p.report_date, None),
    ("Reporting currency", lambda p: p.reporting_currency, None),
    ("Total term debt (mm)", lambda p: p.metrics.get("xbrl_total_debt"), "mm"),
    ("Gross debt incl. CP (mm)", lambda p: p.metrics.get("gross_debt"), "mm"),
    ("Net debt (mm)", lambda p: p.metrics.get("net_debt"), "mm"),
    ("Net debt / EBITDA", lambda p: p.metrics.get("net_debt_to_ebitda"), "x"),
    ("Interest coverage", lambda p: p.metrics.get("interest_coverage"), "x"),
    ("Implied cost of debt", lambda p: p.metrics.get("implied_cost_of_debt_pct"), "pct100"),
    ("Weighted avg coupon", lambda p: p.metrics.get("weighted_avg_coupon_pct"), "pct100"),
    ("Weighted avg maturity (yrs)", lambda p: p.metrics.get("weighted_avg_maturity_yrs"), "yrs"),
    ("Due within 24 months", lambda p: p.metrics.get("due_within_24m_pct"), "pct100"),
    ("Bonds identified", lambda p: p.metrics.get("instruments_identified"), "int"),
    ("Coverage of total debt", lambda p: p.metrics.get("coverage_ratio"), "pct"),
]


def write_comparison(profiles: list[DebtProfile], path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Issuer comparison - debt structure"
    ws["A1"].font = TITLE
    ws["A2"] = ("Values computed by DebtScope from each issuer's latest filing (see per-issuer "
                "workbooks for formulas and audit trail). Mixed reporting currencies are not converted.")
    ws["A2"].font = SUB
    ws.cell(row=4, column=1, value="Metric")
    for c, p in enumerate(profiles, 2):
        ws.cell(row=4, column=c, value=p.ticker or p.company)
    for c in range(1, len(profiles) + 2):
        ws.cell(row=4, column=c).font = H
        ws.cell(row=4, column=c).fill = HEAD_FILL
        ws.cell(row=4, column=c).alignment = Alignment(horizontal="center")
    fmts = {"mm": MM, "x": MULT, "pct": PCT, "pct100": PCT, "yrs": YRS, "int": "0"}
    for r, (label, fn, kind) in enumerate(COMPARE_ROWS, 5):
        ws.cell(row=r, column=1, value=label).font = LBL
        for c, p in enumerate(profiles, 2):
            v = fn(p)
            if v is not None and kind == "mm":
                v = v / 1e6
            if v is not None and kind == "pct100":
                v = v / 100
            cell = ws.cell(row=r, column=c, value="n/a" if v is None else v)
            cell.font = INP
            cell.alignment = Alignment(horizontal="right")
            if kind:
                cell.number_format = fmts[kind]
            if (r % 2) == 0:
                cell.fill = BAND_FILL
        if (r % 2) == 0:
            ws.cell(row=r, column=1).fill = BAND_FILL
    ws.column_dimensions["A"].width = 32
    for c in range(2, len(profiles) + 2):
        ws.column_dimensions[get_column_letter(c)].width = 16
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
