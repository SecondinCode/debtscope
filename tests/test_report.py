import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from debtscope.cli import main

FIX = Path(__file__).parent / "fixtures"


class ReportTests(unittest.TestCase):
    def test_offline_cli_builds_workbook(self):
        with tempfile.TemporaryDirectory() as d:
            code = main(["DEMO", "--offline", str(FIX / "demo_10k.html"),
                         str(FIX / "demo_companyfacts.json"), "--period", "2025-12-31",
                         "--accession", "0000000000-26-000001", "--out", d, "--json"])
            self.assertEqual(code, 0)
            files = sorted(p.name for p in Path(d).iterdir())
            self.assertEqual(len(files), 3)                     # .json, .xlsx, _crawler.html
            xlsx = next(f for f in files if f.endswith(".xlsx"))
            page = (Path(d) / next(f for f in files if f.endswith(".html"))).read_text(encoding="utf-8")
            self.assertIn('data-k="3.600|2042"', page)          # bond mention linked to debt stack
            self.assertIn('id="ds-panel"', page)
            # presentation layer is embedded: walking spider, dashboard strip, panel logic
            for marker in ('id="ds-canvas"', 'id="dsd"', "window.Spider", "window.Dash", "Spider.init("):
                self.assertIn(marker, page)
            wb = load_workbook(Path(d) / xlsx)
            self.assertEqual(wb.sheetnames,
                             ["Summary", "Maturity Wall", "Instruments", "XBRL", "Audit Trail"])
            ws = wb["Summary"]
            self.assertEqual(ws["A14"].value, "Total term debt")
            self.assertTrue(str(ws["C28"].value).startswith("=IFERROR(C27/IF(ISNUMBER("))
            self.assertEqual(len(wb["Maturity Wall"]._charts), 2)


if __name__ == "__main__":
    unittest.main()
